"""Stage 1: inventory a workbook without running it.

Reads the .xlsx or .xlsm file as XML (macros never run, links are never
updated) and records: sheets and their state, used ranges, formula and value
counts, named ranges (with Solver settings), Excel tables, drop-down lists and
other data validation, conditional formats, hidden sheets, rows and columns,
comments, charts, pivot tables, external links, Power Query and other
connections, form controls, VBA presence (module list and code when oletools is
installed), Excel 4 macro sheets, array formulas, what-if data tables and the
calculation settings. Writes one CSV snapshot of the stored values and one of the
formulas per sheet.

Usage:
    python inventory.py <workbook.xlsx|.xlsm> --out <ingest folder>
        [--crosscheck]       compare formula counts with openpyxl (if installed)
        [--excel-snapshots]  PDF of each sheet through Excel (Windows, Excel
                             installed). Not yet tested. Starts its own hidden
                             Excel instance with macros disabled and links not
                             updated; never touches workbooks you have open.

Writes inventory.json, snapshots/*.csv and, with oletools, vba/*.bas.
"""

import argparse
import csv
import hashlib
import json
import re
import sys
from collections import Counter
from pathlib import Path

import xlsxlib as X


def _safe(name):
    return re.sub(r"[^\w.-]+", "_", name).strip("_") or "sheet"


def used_range(sheet):
    if not sheet.cells:
        return None
    rows = [r for r, _ in sheet.cells]
    cols = [c for _, c in sheet.cells]
    return X.range_text(min(rows), min(cols), max(rows), max(cols))


def name_kind(wb, item):
    text = item.text.strip()
    if item.name.lower().startswith("solver_"):
        return "solver"
    if item.name.startswith("_xlnm."):
        return "builtin"
    if X.defined_name_ranges(wb, item):
        return "range"
    if re.fullmatch(r"=?-?[\d.]+(E[+-]?\d+)?|=?\".*\"|=?TRUE|=?FALSE", text, re.I):
        return "constant"
    if "LAMBDA(" in text.upper():
        return "lambda"
    return "formula"


def read_vba(path):
    """Module names, procedures and code through oletools (optional)."""
    try:
        from oletools.olevba import VBA_Parser
    except ImportError:
        return {"available": False,
                "message": "VBA is present. oletools is not installed, so the modules were "
                           "not read (pip install oletools). Ask the owner for the code."}
    modules = []
    parser = VBA_Parser(str(path))
    try:
        for _, _, name, code in parser.extract_macros():
            procs = re.findall(r"^\s*(?:Public\s+|Private\s+)?(Sub|Function)\s+(\w+)", code,
                               re.M | re.I)
            modules.append({"module": name, "code": code,
                            "procedures": [{"kind": k.title(), "name": n} for k, n in procs]})
    finally:
        parser.close()
    autorun = sorted({p["name"] for m in modules for p in m["procedures"]
                      if p["name"].lower() in ("workbook_open", "auto_open", "workbook_activate",
                                               "workbook_beforeclose", "auto_close")})
    return {"available": True, "modules": modules, "autorun": autorun}


def excel_snapshots(path, folder):
    """PDF snapshot of each visible worksheet through Excel COM. Not yet tested.

    A new, hidden Excel instance is started (DispatchEx), so workbooks you have
    open in Excel are not touched. Macros are disabled before the file opens
    (AutomationSecurity = 3, force disable), links are not updated
    (UpdateLinks = 0) and the file opens read-only. Nothing is saved."""
    try:
        import win32com.client
    except ImportError:
        return "pywin32 is not installed; Excel snapshots skipped."
    folder.mkdir(parents=True, exist_ok=True)
    excel = win32com.client.DispatchEx("Excel.Application")
    try:
        excel.Visible = False
        excel.DisplayAlerts = False
        excel.ScreenUpdating = False
        excel.EnableEvents = False
        excel.AutomationSecurity = 3
        book = excel.Workbooks.Open(str(Path(path).resolve()), UpdateLinks=0, ReadOnly=True,
                                    IgnoreReadOnlyRecommended=True, AddToMru=False,
                                    Notify=False)
        written = 0
        for sheet in book.Worksheets:
            if sheet.Visible != -1:
                continue
            target = folder / f"{sheet.Index:02d}-{_safe(sheet.Name)}.pdf"
            sheet.ExportAsFixedFormat(0, str(target))
            written += 1
        book.Close(SaveChanges=False)
        return f"{written} PDF snapshots written to {folder}"
    finally:
        excel.Quit()


def crosscheck(path, wb):
    """Formula counts per sheet from openpyxl next to this reader's counts."""
    try:
        import warnings
        import openpyxl
    except ImportError:
        return {"available": False, "message": "openpyxl is not installed; cross-check skipped."}
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        book = openpyxl.load_workbook(path, read_only=True)
    rows = []
    for ws in book.worksheets:
        count = 0
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and value.startswith("="):
                    count += 1
                elif value is not None and type(value).__name__ in ("ArrayFormula",
                                                                     "DataTableFormula"):
                    count += 1
        mine = wb.sheet(ws.title)
        own = 0
        if mine is not None:
            for c in mine.cells.values():
                member = c.ftype == "array-member" or (c.ftype == "dataTable" and not c.fref)
                if c.formula is not None and not member:
                    own += 1
        rows.append({"sheet": ws.title, "openpyxl": count, "reader": own,
                     "match": count == own})
    book.close()
    return {"available": True, "sheets": rows}


def build(wb, path):
    raw = Path(path).read_bytes()
    functions = Counter()
    sheets = []
    for sheet in wb.sheets:
        formulas = [c for c in sheet.cells.values() if c.formula is not None]
        for cell in formulas:
            if cell.ftype in ("array-member",):
                continue
            for tok in X.tokenize(cell.formula):
                if tok.kind == "func":
                    functions[tok.value] += 1
        values = [c for c in sheet.cells.values() if c.formula is None and c.value is not None]
        sheets.append({
            "name": sheet.name, "index": sheet.index, "state": sheet.state, "kind": sheet.kind,
            "usedRange": used_range(sheet), "cells": len(sheet.cells),
            "formulas": len(formulas),
            "numbers": sum(1 for c in values if not c.is_text and c.vtype != "e"),
            "texts": sum(1 for c in values if c.is_text),
            "errors": sum(1 for c in sheet.cells.values() if c.vtype == "e"),
            "hiddenRows": sorted(sheet.hidden_rows)[:200],
            "hiddenColumns": [X.num_to_col(c) for c in sorted(sheet.hidden_cols)][:200],
            "merged": sheet.merged[:200], "protected": sheet.protected,
            "comments": [{"cell": k, "text": v} for k, v in list(sheet.comments.items())[:200]],
            "tables": sheet.tables, "charts": sheet.charts, "pivots": sheet.pivots,
            "formControls": sheet.controls, "images": sheet.has_drawing_images,
            "dataValidations": sheet.validations,
            "conditionalFormats": sheet.conditional_formats,
        })
    names = [{"name": n.name, "scope": None if n.scope is None else wb.sheets[n.scope].name,
              "hidden": n.hidden, "refersTo": "=" + n.text.lstrip("="),
              "kind": name_kind(wb, n)} for n in wb.names]
    solver = {n["name"]: n["refersTo"] for n in names if n["kind"] == "solver"}
    return {
        "file": {"name": Path(path).name, "bytes": len(raw),
                 "sha256": hashlib.sha256(raw).hexdigest(),
                 "format": Path(path).suffix.lower().lstrip("."),
                 "date1904": wb.date1904, "workbookProtected": wb.protected},
        "calculation": {"settings": wb.calc,
                        "mode": wb.calc.get("calcMode", "auto"),
                        "iterative": wb.calc.get("iterate") in ("1", "true")},
        "sheets": sheets,
        "definedNames": names,
        "solver": solver,
        "externalLinks": wb.external_links,
        "connections": wb.connections,
        "powerQuery": wb.custom_xml_mashup or any(c["powerQuery"] for c in wb.connections),
        "queryTables": wb.query_tables,
        "pivotCaches": wb.pivot_caches,
        "vbaPresent": wb.vba,
        "macroSheets": wb.macrosheets,
        "arrayFormulas": wb.array_formulas[:500],
        "dataTables": wb.data_tables,
        "functions": dict(functions.most_common()),
        "totals": {"sheets": len(wb.sheets),
                   "formulas": sum(s["formulas"] for s in sheets),
                   "numbers": sum(s["numbers"] for s in sheets),
                   "texts": sum(s["texts"] for s in sheets),
                   "errors": sum(s["errors"] for s in sheets),
                   "charts": sum(len(s["charts"]) for s in sheets),
                   "pivots": sum(len(s["pivots"]) for s in sheets),
                   "tables": sum(len(s["tables"]) for s in sheets),
                   "dataValidations": sum(len(s["dataValidations"]) for s in sheets),
                   "conditionalFormats": sum(len(s["conditionalFormats"]) for s in sheets)},
        "notes": wb.notes,
    }


def write_snapshots(wb, folder):
    folder.mkdir(parents=True, exist_ok=True)
    for sheet in wb.sheets:
        if not sheet.cells:
            continue
        rows = [r for r, _ in sheet.cells]
        cols = [c for _, c in sheet.cells]
        r1, r2, c1, c2 = min(rows), min(max(rows), min(rows) + 4999), min(cols), max(cols)
        stem = f"{sheet.index + 1:02d}-{_safe(sheet.name)}"
        for suffix, pick in (("values", lambda c: c.value),
                             ("formulas", lambda c: "=" + c.formula if c.formula else c.value)):
            with open(folder / f"{stem}.{suffix}.csv", "w", newline="",
                      encoding="utf-8") as fh:
                writer = csv.writer(fh)
                writer.writerow([""] + [X.num_to_col(c) for c in range(c1, c2 + 1)])
                for r in range(r1, r2 + 1):
                    line = [r]
                    for c in range(c1, c2 + 1):
                        cell = sheet.get(r, c)
                        value = None if cell is None else pick(cell)
                        line.append("" if value is None else value)
                    writer.writerow(line)


def run(path, out, cross=False, snapshots_pdf=False):
    out = Path(out)
    out.mkdir(parents=True, exist_ok=True)
    wb = X.load(path)
    data = build(wb, path)
    if wb.vba:
        vba = read_vba(path)
        if vba.get("available"):
            vba_dir = out / "vba"
            vba_dir.mkdir(exist_ok=True)
            for module in vba["modules"]:
                (vba_dir / f"{_safe(module['module'])}.bas").write_text(module["code"],
                                                                        encoding="utf-8")
                module["code"] = f"vba/{_safe(module['module'])}.bas"
        data["vba"] = vba
    if cross:
        data["crosscheck"] = crosscheck(path, wb)
    write_snapshots(wb, out / "snapshots")
    if snapshots_pdf:
        data["excelSnapshots"] = excel_snapshots(path, out / "snapshots")
    (out / "inventory.json").write_text(json.dumps(data, indent=1, default=str),
                                        encoding="utf-8")
    return wb, data


def main():
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    parser.add_argument("workbook")
    parser.add_argument("--out", required=True, help="ingest folder for the results")
    parser.add_argument("--crosscheck", action="store_true")
    parser.add_argument("--excel-snapshots", action="store_true")
    args = parser.parse_args()
    _, data = run(args.workbook, args.out, args.crosscheck, args.excel_snapshots)
    t = data["totals"]
    print(f"{t['sheets']} sheets, {t['formulas']} formulas, {t['numbers']} numbers, "
          f"{t['errors']} errors, {t['charts']} charts, {len(data['externalLinks'])} external "
          f"links, VBA {'yes' if data['vbaPresent'] else 'no'} -> "
          f"{Path(args.out) / 'inventory.json'}")
    if data.get("vba") and not data["vba"].get("available"):
        print(data["vba"]["message"])
    return 0


if __name__ == "__main__":
    sys.exit(main())
